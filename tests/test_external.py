"""Внешние источники: чужие сайты с брендом из ответов ИИ — выгрузка и зачёт.

Сеть не нужна: проверка страниц подменена. Проверяется то, за что отвечает
модуль: сайт клиента не считается внешним, страница проверяется один раз,
находка превращает «не найдено» в «найдено» с типом `source`, а Excel —
ровно два столбца «Источник» и «Текст упоминания».
"""

import io
import json
import sys
import tempfile
import traceback
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

_tmp_db = Path(tempfile.gettempdir()) / "aiparser_test_external.db"
_tmp_db.unlink(missing_ok=True)

from app import config  # noqa: E402

config.DB_PATH = _tmp_db

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.api import create_app  # noqa: E402
from app.db import repo  # noqa: E402
from app.detect import deep, external  # noqa: E402

client = TestClient(create_app())

PAGES = {
    "https://partner.ru/catalog/geosoft": (True, "…в каталоге представлена продукция Геософт Дент…", None),
    "https://other.ru/a": (False, None, None),
    "https://shop.ru/z": (False, None, "ConnectError"),
}
calls: list[list[str]] = []


async def fake_check_pages(urls, brand_name, aliases, **_):
    calls.append(list(urls))
    return [deep.PageCheck(u, *PAGES[u]) for u in urls]


external.deep.check_pages = fake_check_pages

_n = 0


def _setup() -> tuple[int, dict[str, int]]:
    global _n
    _n += 1
    p = client.post("/api/projects", json={
        "name": f"Геософт {_n}", "brand_name": "Геософт Дент", "brand_aliases": ["Geosoft"],
        "brand_domains": ["geosoft-dent.ru"],
    }).json()
    client.post(f"/api/projects/{p['id']}/queries", json={"text": "стоматологическое оборудование\nГеософт отзывы"})
    q = {x["text"]: x["id"] for x in client.get(f"/api/projects/{p['id']}/queries").json()}
    sid = repo.create_scan(p["id"], ["alice", "google_aio"], {})
    repo._exec("UPDATE scans SET scan_date = '2026-09-11' WHERE id = ?", (sid,))
    repo.save_result(sid, q["стоматологическое оборудование"], "alice", "not_found",
                     sources=["https://geosoft-dent.ru/x", "https://partner.ru/catalog/geosoft", "https://other.ru/a"])
    repo.save_result(sid, q["Геософт отзывы"], "google_aio", "found", mention_types=["text"],
                     evidence_quote="Геософт Дент — производитель", detected_by="rules", confidence=0.95,
                     sources=["https://partner.ru/catalog/geosoft", "https://shop.ru/z"])
    repo.save_result(sid, q["Геософт отзывы"], "alice", "error", sources=["https://partner.ru/catalog/geosoft"])
    return p["id"], q


def _row(pid: int, qid: int, service: str) -> dict:
    return next(r for r in repo.results_with_sources_on_date(pid, "2026-09-11")
                if r["query_id"] == qid and r["service"] == service)


def test_check_counts_source_mentions_and_skips_own_site():
    calls.clear()
    pid, q = _setup()
    r = client.post(f"/api/projects/{pid}/external-sources/check").json()
    assert r["date"] == "2026-09-11"
    # сайт клиента не проверяется и не считается внешним
    assert sorted(calls[0]) == ["https://other.ru/a", "https://partner.ru/catalog/geosoft", "https://shop.ru/z"]
    assert r["urls_total"] == 3 and r["found_sites"] == 1 and r["failed"] == 1
    assert r["sites"] == [{"url": "https://partner.ru/catalog/geosoft", "quote": PAGES["https://partner.ru/catalog/geosoft"][1]}]

    # «не найдено» стало «найдено» на сайте-источнике
    a = _row(pid, q["стоматологическое оборудование"], "alice")
    assert a["status"] == "found" and json.loads(a["mention_types_json"]) == ["source"]
    assert "partner.ru" in a["evidence_quote"] and a["detected_by"] == "deep"
    # у найденного добавился тип, цитата прежняя
    g = _row(pid, q["Геософт отзывы"], "google_aio")
    assert json.loads(g["mention_types_json"]) == ["source", "text"]
    assert g["evidence_quote"] == "Геософт Дент — производитель"
    # ошибка проверки так и остаётся ошибкой — ответа не было
    assert _row(pid, q["Геософт отзывы"], "alice")["status"] == "error"
    assert r["updated_results"] == 2


def test_pages_are_checked_once():
    calls.clear()
    pid, _ = _setup()
    client.post(f"/api/projects/{pid}/external-sources/check")
    again = client.post(f"/api/projects/{pid}/external-sources/check").json()
    assert len(calls) == 1, "повторная проверка снова пошла в сеть"
    assert again["checked_now"] == 0 and again["updated_results"] == 0


def test_brand_change_invalidates_cache():
    calls.clear()
    pid, _ = _setup()
    client.post(f"/api/projects/{pid}/external-sources/check")
    client.patch(f"/api/projects/{pid}", json={"brand_aliases": ["Geosoft", "Геософт"]})
    client.post(f"/api/projects/{pid}/external-sources/check")
    assert len(calls) == 2


def test_source_mention_is_withdrawn_when_site_no_longer_matches():
    # Сначала партнёрский сайт с брендом засчитан, потом (сменили алиасы или
    # правила) бренд на нём больше не находится — «найдено» должно откатиться.
    pid, q = _setup()
    client.post(f"/api/projects/{pid}/external-sources/check")
    assert _row(pid, q["стоматологическое оборудование"], "alice")["status"] == "found"

    saved = PAGES["https://partner.ru/catalog/geosoft"]
    PAGES["https://partner.ru/catalog/geosoft"] = (False, None, None)
    try:
        client.patch(f"/api/projects/{pid}", json={"brand_aliases": ["Geosoft", "ГеоСофт"]})  # сброс кэша
        r = client.post(f"/api/projects/{pid}/external-sources/check").json()
    finally:
        PAGES["https://partner.ru/catalog/geosoft"] = saved

    a = _row(pid, q["стоматологическое оборудование"], "alice")
    assert a["status"] == "not_found" and json.loads(a["mention_types_json"]) == []
    g = _row(pid, q["Геософт отзывы"], "google_aio")
    assert g["status"] == "found" and json.loads(g["mention_types_json"]) == ["text"]
    assert r["found_sites"] == 0 and r["updated_results"] == 2


def test_page_matching_is_strict():
    # Целым словом, без нечёткого сравнения, без <head>, с раскодированием.
    page = deep.html_to_text(
        "<html><head><title>Neighbors App</title></head><body>"
        "<p>Neighborbrite Landscape Design</p><p>Магазин &quot;Геософт Дент&quot; в каталоге</p></body></html>"
    )
    assert "Neighbors App" not in page and '"Геософт Дент"' in page
    assert not deep.check_page_text("Neighborbrite Landscape Design", "Neighbors", []).found
    v = deep.check_page_text(page, "Геософт Дент", [])
    assert v.found and "Геософт Дент" in (v.evidence_quote or "")


def test_xlsx_has_two_columns():
    pid, _ = _setup()
    resp = client.get(f"/api/projects/{pid}/external-sources.xlsx", params={"date": "2026-09-11"})
    assert resp.status_code == 200, resp.text
    assert "spreadsheetml" in resp.headers["content-type"]
    assert "attachment" in resp.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(resp.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Источник", "Текст упоминания")
    assert rows[1:] == [("https://partner.ru/catalog/geosoft", PAGES["https://partner.ru/catalog/geosoft"][1])]


def test_bad_date_and_unknown_project():
    pid, _ = _setup()
    assert client.post(f"/api/projects/{pid}/external-sources/check", params={"date": "11.09.2026"}).status_code == 400
    assert client.post("/api/projects/999999/external-sources/check").status_code == 404


def test_is_external():
    d = ["geosoft-dent.ru"]
    assert not external.is_external("https://geosoft-dent.ru/a", d)
    assert not external.is_external("https://shop.geosoft-dent.ru/a", d)
    assert external.is_external("https://partner.ru/a", d)
    assert not external.is_external("mailto:a@b.ru", d)


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    failed = 0
    for name, fn in tests:
        try:
            fn()
        except Exception:
            failed += 1
            print(f"FAIL {name}")
            traceback.print_exc()
    print(f"{len(tests) - failed} passed, {failed} failed")
    sys.exit(1 if failed else 0)
