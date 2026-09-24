"""Выгрузка упоминаемости: запросы в строках, ИИ-системы в столбцах.

Файл должен повторять таблицу на экране: те же знаки в ячейках, тот же период
из календаря и те же системы. Проверяется структура книги, а не вёрстка.
"""

import io
import sys
import tempfile
import traceback
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

_tmp_db = Path(tempfile.gettempdir()) / "aiparser_test_mentions.db"
_tmp_db.unlink(missing_ok=True)

from app import config  # noqa: E402

config.DB_PATH = _tmp_db

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.api import create_app  # noqa: E402
from app.db import repo  # noqa: E402

client = TestClient(create_app())
_n = 0


def _project(queries: list[str]) -> tuple[int, list[int]]:
    global _n
    _n += 1
    pid = repo.create_project(f"Упоминаемость {_n}", "Геософт")
    repo.add_queries(pid, queries)
    return pid, [q["id"] for q in repo.list_queries(pid)]


def _scan(pid: int, day: str, rows: dict[str, dict[int, str]]) -> None:
    """rows: сервис → {query_id: статус}."""
    sid = repo.create_scan(pid, list(rows), {})
    repo._exec("UPDATE scans SET scan_date = ?, status = 'done' WHERE id = ?", (day, sid))
    for service, by_query in rows.items():
        for qid, status in by_query.items():
            repo.save_result(sid, qid, service, status)


def _book(pid: int, params: str = ""):
    r = client.get(f"/api/projects/{pid}/mentions.xlsx?{params}")
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    assert "attachment" in r.headers["content-disposition"]
    return load_workbook(io.BytesIO(r.content))


def test_one_date_is_a_flat_table():
    pid, q = _project(["первый запрос", "второй запрос"])
    _scan(pid, "2026-09-14", {
        "perplexity": {q[0]: "found", q[1]: "not_found"},
        "alice": {q[0]: "error"},
    })
    ws = _book(pid)["Упоминаемость"]
    assert [c.value for c in ws[1]] == ["Запрос", "Perplexity", "Алиса AI"]
    assert [c.value for c in ws[2]] == ["первый запрос", "✓", "!"]
    # Алиса этот запрос в тот день не проверяла — не «нет упоминаний», а «нет данных».
    assert [c.value for c in ws[3]] == ["второй запрос", "✗", "·"]
    assert ws.max_row == 3


def test_period_groups_columns_by_date():
    pid, q = _project(["запрос"])
    _scan(pid, "2026-09-10", {"perplexity": {q[0]: "found"}, "alice": {q[0]: "not_found"}})
    _scan(pid, "2026-09-14", {"perplexity": {q[0]: "not_found"}, "alice": {q[0]: "found"}})
    ws = _book(pid, "mode=period")["Упоминаемость"]
    # Дата стоит над своей группой систем: ячейки группы объединены, поэтому
    # у соседних колонок значение пустое.
    assert [c.value for c in ws[1]] == ["Запрос", "2026-09-10", None, "2026-09-14", None]
    assert [c.value for c in ws[2]] == [None, "Perplexity", "Алиса AI", "Perplexity", "Алиса AI"]
    assert [c.value for c in ws[3]] == ["запрос", "✓", "✗", "✗", "✓"]


def test_calendar_range_narrows_the_export():
    pid, q = _project(["запрос"])
    _scan(pid, "2026-09-10", {"perplexity": {q[0]: "found"}})
    _scan(pid, "2026-09-14", {"perplexity": {q[0]: "not_found"}})
    ws = _book(pid, "mode=period&date_from=2026-09-14&date_to=2026-09-14")["Упоминаемость"]
    assert [c.value for c in ws[1]] == ["Запрос", "Perplexity"]
    assert [c.value for c in ws[2]] == ["запрос", "✗"]


def test_legend_sheet_explains_the_signs():
    pid, q = _project(["запрос"])
    _scan(pid, "2026-09-14", {"perplexity": {q[0]: "limit_reached"}})
    wb = _book(pid)
    assert wb["Упоминаемость"]["B2"].value == "◷"
    legend = {row[0]: str(row[1]) for row in wb["Обозначения"].iter_rows(min_row=2, values_only=True) if row[0]}
    assert legend["◷"].startswith("Лимит тарифа")
    assert legend["·"] == "Не проверялся в этот день"


def test_errors_are_explicit():
    pid, _ = _project(["запрос"])
    # Сканов ещё не было — выгружать нечего, и это не пустой файл, а понятный отказ.
    assert client.get(f"/api/projects/{pid}/mentions.xlsx").status_code == 404
    assert client.get("/api/projects/99999/mentions.xlsx").status_code == 404
    assert client.get(f"/api/projects/{pid}/mentions.xlsx?mode=ерунда").status_code == 400


def test_product_card_toggle_matches_dashboard_and_export():
    pid, q = _project(["только карточка", "карточка и текст"])
    sid = repo.create_scan(pid, ["perplexity"], {})
    repo._exec("UPDATE scans SET scan_date = ?, status = 'done' WHERE id = ?", ("2026-09-15", sid))
    repo.save_result(sid, q[0], "perplexity", "found", mention_types=["card"])
    repo.save_result(sid, q[1], "perplexity", "found", mention_types=["card", "text"])

    included = client.get(f"/api/projects/{pid}/overview?include_cards=true").json()
    excluded = client.get(f"/api/projects/{pid}/overview?include_cards=false").json()
    assert included["summary"]["total"]["found"] == 2
    assert excluded["summary"]["total"]["found"] == 1
    assert excluded["summary"]["total"]["checked"] == 2
    assert excluded["selection"]["include_cards"] is False

    ws = _book(pid, "include_cards=false")["Упоминаемость"]
    assert [c.value for c in ws[2]] == ["только карточка", "✗"]
    assert [c.value for c in ws[3]] == ["карточка и текст", "✓"]


if __name__ == "__main__":
    tests = [(n, f) for n, f in list(globals().items()) if n.startswith("test_") and callable(f)]
    failed = 0
    for name, fn in tests:
        try:
            fn()
            print(f"OK   {name}")
        except Exception:
            print(f"FAIL {name}")
            traceback.print_exc()
            failed += 1
    print(f"\n{len(tests) - failed} passed, {failed} failed")
    sys.exit(1 if failed else 0)
