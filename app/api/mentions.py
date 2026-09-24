"""Выгрузка упоминаемости: запросы в строках, ИИ-системы в столбцах.

Тот же срез, что видно в таблице на дашборде, и те же знаки в ячейках —
человек, открывший файл, должен узнать в нём экран программы, а не гадать,
почему цифры разошлись. Даты берутся из того же календаря: за период
получается по группе столбцов на каждую дату.
"""

from __future__ import annotations

import io
import re
from collections import defaultdict
from urllib.parse import quote

from fastapi import APIRouter, HTTPException
from fastapi.responses import Response

from app import services
from app.api.results import MAX_DATES, MODES, SCOPES, _check_date, allowed_mention_types, scoped_status, select_dates
from app.db import repo

router = APIRouter(prefix="/api", tags=["mentions"])

# Знаки — ровно те же, что в ячейках таблицы (frontend/src/lib/status.ts).
SIGNS = {
    "found": "✓",
    "not_found": "✗",
    "skipped": "—",
    "error": "!",
    "auth_required": "!",
    "captcha": "!",
    "limit_reached": "◷",
}
# Запрос по этой системе в этот день не проверялся.
NO_DATA = "·"

LEGEND = [
    ("✓", "Упоминание найдено"),
    ("✗", "Упоминаний нет"),
    ("—", "AI-блок не показан"),
    ("!", "Проверка не удалась: ошибка, капча или нужен вход"),
    ("◷", "Лимит тарифа — запрос не проверен"),
    (NO_DATA, "Не проверялся в этот день"),
]


def collect(project_id: int, dates: list[str], allowed: set[str] | None = None) -> tuple[list[dict], list[str]]:
    """Строки «запрос → дата → сервис → знак» и список сервисов с данными.

    `allowed` — учёт упоминаний с экрана (см. SCOPES): файл обязан совпадать
    с таблицей, иначе выгрузка и дашборд разойдутся в цифрах.
    """
    results = repo.results_by_date(project_id, dates)
    cells: dict[int, dict[str, dict[str, str]]] = defaultdict(lambda: defaultdict(dict))
    with_data: set[str] = set()
    for r in results:
        status = scoped_status(r["status"], r.get("mention_types_json"), allowed)
        cells[r["query_id"]][r["scan_date"]][r["service"]] = SIGNS.get(status, "?")
        with_data.add(r["service"])

    # Порядок систем — как в интерфейсе, а не как в базе; системы без единой
    # проверки в выбранном периоде не показываем, иначе таблица зарастёт
    # пустыми столбцами.
    service_ids = [s.id for s in services.SERVICES if s.id in with_data]

    rows = []
    for q in repo.list_queries(project_id):
        qc = cells.get(q["id"])
        if not q["is_active"] and not qc:
            continue
        rows.append({"text": q["text"], "cells": qc or {}})
    return rows, service_ids


def build_xlsx(project: dict, dates: list[str], service_ids: list[str], rows: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    names = {s.id: s.name for s in services.SERVICES}
    wb = Workbook()
    ws = wb.active
    ws.title = "Упоминаемость"

    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    head_fill = PatternFill("solid", fgColor="EEF2F7")
    many = len(dates) > 1

    if many:
        # Две строки шапки: сверху дата, под ней системы этой даты.
        ws.append(["Запрос"] + [d for d in dates for _ in service_ids])
        ws.append([""] + [names.get(s, s) for _ in dates for s in service_ids])
        ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
        for i, _ in enumerate(dates):
            first = 2 + i * len(service_ids)
            ws.merge_cells(start_row=1, start_column=first, end_row=1,
                           end_column=first + len(service_ids) - 1)
        head_rows, first_data_row = (1, 2), 3
    else:
        ws.append(["Запрос"] + [names.get(s, s) for s in service_ids])
        head_rows, first_data_row = (1,), 2

    for r in head_rows:
        for cell in ws[r]:
            cell.font = bold
            cell.alignment = center
            cell.fill = head_fill

    for row in rows:
        line = [row["text"]]
        for d in dates:
            by_service = row["cells"].get(d, {})
            line += [by_service.get(s, NO_DATA) for s in service_ids]
        ws.append(line)

    ws.column_dimensions["A"].width = 70
    for i in range(len(dates) * len(service_ids) if many else len(service_ids)):
        ws.column_dimensions[get_column_letter(2 + i)].width = 16
    for line in ws.iter_rows(min_row=first_data_row):
        line[0].alignment = Alignment(wrap_text=True, vertical="top")
        for cell in line[1:]:
            cell.alignment = center
    ws.freeze_panes = ws.cell(row=first_data_row, column=2)

    legend = wb.create_sheet("Обозначения")
    legend.append(["Знак", "Что значит"])
    for cell in legend[1]:
        cell.font = bold
    for sign, text in LEGEND:
        legend.append([sign, text])
    legend.column_dimensions["A"].width = 8
    legend.column_dimensions["B"].width = 60
    legend.append([])
    legend.append(["Проект", project["name"]])
    legend.append(["Бренд", project["brand_name"]])
    legend.append(["Даты", ", ".join(dates)])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/projects/{project_id}/mentions.xlsx")
def export_mentions(
    project_id: int,
    days: int = 30,
    mode: str = "period",
    date_from: str | None = None,
    date_to: str | None = None,
    dates: str | None = None,
    max_dates: int = MAX_DATES,
    scope: str = "all",
    include_cards: bool = True,
) -> Response:
    """Excel «запросы × ИИ-системы» за выбранный в календаре период."""
    project = repo.get_project(project_id)
    if not project:
        raise HTTPException(404, "Проект не найден")
    if mode not in MODES:
        raise HTTPException(400, f"mode: одно из {', '.join(MODES)}")
    if scope not in SCOPES:
        raise HTTPException(400, f"scope: одно из {', '.join(SCOPES)}")
    date_from = _check_date(date_from, "date_from")
    date_to = _check_date(date_to, "date_to")
    picked = [d.strip() for d in (dates or "").split(",") if d.strip()]
    for d in picked:
        _check_date(d, "dates")

    all_dates = [r["date"] for r in repo.scan_dates(project_id)]
    if not all_dates:
        raise HTTPException(404, "По проекту ещё не было сканов")
    selected, _available = select_dates(
        all_dates,
        mode=mode,
        date_from=date_from,
        date_to=date_to,
        picked=picked,
        days=min(days, 365),
        max_dates=min(max(1, max_dates), MAX_DATES),
    )
    if not selected:
        raise HTTPException(404, "За выбранный период нет ни одной проверки")

    rows, service_ids = collect(project_id, selected, allowed_mention_types(scope, include_cards))
    body = build_xlsx(project, selected, service_ids, rows)

    period = selected[0] if len(selected) == 1 else f"{selected[0]}—{selected[-1]}"
    name = re.sub(r"[^\w.-]+", "-", f"упоминаемость-{project['name']}-{period}").strip("-") + ".xlsx"
    return Response(
        content=body,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"},
    )
