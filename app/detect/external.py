"""Внешние источники: чужие сайты, на которые ссылался ИИ и где упомянут бренд.

Зачем. Клиент может влиять на сайты, где публикуют его продукцию (дилеры,
каталоги, партнёры). Если ИИ цитирует такой сайт, а на нём есть бренд, — это
тоже упоминание: пользователь ИИ дойдёт до бренда через источник. Отсюда две
вещи:

* список таких сайтов за срез — выгрузка в Excel «источник + текст
  упоминания», чтобы работать с ними;
* сами результаты: ответ, который процитировал такой сайт, получает тип
  упоминания `source` («на сайте-источнике»), а «не найдено» становится
  «найдено».

Каждая страница проверяется один раз на проект (таблица source_pages): в срезе
сотни ссылок, и многие сайты повторяются из запроса в запрос. Кэш сбрасывается
сам, если поменялись формы бренда.
"""

from __future__ import annotations

import io
import json
import logging
from urllib.parse import urlparse

from app.db import repo
from app.detect import deep
from app.detect.brand import normalize_host, same_site

log = logging.getLogger("aiparser.external")

SOURCE = "source"
# Версия правил проверки страниц. Поменялись правила — старый кэш не годится.
# 2 (11.09.2026): бренд на странице только целым словом, без нечёткого
# сравнения, без <head>, с раскодированными HTML-сущностями.
MATCHER_VERSION = 2


def brand_sig(project: dict) -> str:
    """Формы бренда и версия правил, при которых страница проверялась."""
    return json.dumps(
        [MATCHER_VERSION, project["brand_name"].strip().lower(),
         sorted(a.strip().lower() for a in project["brand_aliases"])],
        ensure_ascii=False,
    )


def is_external(url: str, brand_domains: list[str]) -> bool:
    """Чужой сайт: http(s)-страница не на домене клиента."""
    try:
        parsed = urlparse(url)
    except ValueError:
        return False
    if parsed.scheme not in ("http", "https"):
        return False
    host = normalize_host(parsed.hostname or "")
    return bool(host) and not any(same_site(host, d) for d in brand_domains)


async def check_date(project: dict, scan_date: str) -> dict:
    """Проверяет все внешние источники среза и засчитывает находки как упоминания."""
    rows = repo.results_with_sources_on_date(project["id"], scan_date)
    domains = project["brand_domains"]

    urls: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for u in json.loads(r["sources_json"] or "[]"):
            if u not in seen and is_external(u, domains):
                seen.add(u)
                urls.append(u)

    sig = brand_sig(project)
    cached = repo.source_pages_get(project["id"], sig)
    todo = [u for u in urls if u not in cached]
    if todo:
        log.info("Внешние источники %s: проверяю %s новых страниц из %s", scan_date, len(todo), len(urls))
        checks = await deep.check_pages(todo, project["brand_name"], project["brand_aliases"])
        repo.source_pages_save(project["id"], sig, checks)
        cached = repo.source_pages_get(project["id"], sig)

    found = {u: cached[u]["quote"] for u in urls if u in cached and cached[u]["found"]}
    updated = _apply_to_results(rows, found)

    return {
        "date": scan_date,
        "urls_total": len(urls),
        "checked_now": len(todo),
        "failed": sum(1 for u in urls if u in cached and cached[u]["error"]),
        "found_sites": len(found),
        "updated_results": updated,
        "sites": [{"url": u, "quote": q or ""} for u, q in sorted(found.items(), key=_site_order)],
    }


def _site_order(item: tuple[str, str | None]) -> tuple[str, str]:
    host = normalize_host(urlparse(item[0]).hostname or "")
    return host, item[0]


def _apply_to_results(rows: list[dict], found: dict[str, str | None]) -> int:
    """Приводит тип `source` в результатах в соответствие с проверкой сайтов.

    В обе стороны: ответ, процитировавший сайт с брендом, получает `source`
    (а «не найдено» становится «найдено»), а у ответа, чьи сайты больше не
    проходят проверку (поменялись формы бренда или правила), `source`
    снимается — иначе однажды засчитанное ложное упоминание осталось бы в
    статистике навсегда.
    """
    updated = 0
    for r in rows:
        if r["status"] not in ("found", "not_found"):
            continue  # сбой или лимит — ответа нет, засчитывать нечего
        hits = [u for u in json.loads(r["sources_json"] or "[]") if u in found]
        types = json.loads(r["mention_types_json"] or "[]")

        if hits:
            if SOURCE in types and r["status"] == "found":
                continue  # уже засчитано
            new_types = sorted(set(types) | {SOURCE})
            if r["status"] == "not_found":
                url = hits[0]
                repo.update_result_mention(
                    r["id"], status="found", mention_types=new_types,
                    evidence_quote=f"На сайте-источнике {url}: {found[url] or ''}".strip(),
                    detected_by="deep", confidence=0.9,
                )
            else:
                repo.update_result_mention(
                    r["id"], status="found", mention_types=new_types,
                    evidence_quote=r["evidence_quote"], detected_by=r["detected_by"], confidence=r["confidence"],
                )
            updated += 1
        elif SOURCE in types:
            rest = [t for t in types if t != SOURCE]
            if rest:
                repo.update_result_mention(
                    r["id"], status="found", mention_types=rest,
                    evidence_quote=r["evidence_quote"], detected_by=r["detected_by"], confidence=r["confidence"],
                )
            else:
                # «Найдено» держалось только на сайте-источнике — его больше нет.
                repo.update_result_mention(
                    r["id"], status="not_found", mention_types=[],
                    evidence_quote=None, detected_by="none", confidence=None,
                )
            updated += 1
    return updated


def build_xlsx(sites: list[dict]) -> bytes:
    """Excel в два столбца: источник и текст упоминания на нём."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Внешние источники"
    ws.append(["Источник", "Текст упоминания"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for s in sites:
        ws.append([s["url"], s["quote"]])
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
