import io
import json
from datetime import date, datetime, timezone
from types import SimpleNamespace
import pytest
from sqlalchemy import event, select

from openpyxl import load_workbook
from server.test_control import setup
from server.models import CloudResult, ControlProject, ProjectMember, User, make_session_factory


def seed_report(tmp_path):
    client, owner, other, devices, url = setup(tmp_path)
    project = client.post('/api/v1/control/projects', headers=owner, json={
        'name': 'geosoft-dent.ru', 'brand_name': 'ESTUS',
        'queries': [{'text': '=Original equipment?', 'group_tag': 'Brand'}, {'text': 'Dental tools'}, {'text': 'Not checked', 'group_tag': 'Brand'}],
        'config': {'services': ['google_aio', 'chatgpt'], 'brand_domains': ['www.geosoft-dent.ru']},
    }).json()
    _, sessions = make_session_factory(url)
    ids = []
    with sessions() as db:
        user_id = db.get(ControlProject, project['id']).user_id
        for i, (query, day, service, status, kinds) in enumerate([
            (0, '2026-09-01', 'google_aio', 'not_found', []),
            (0, '2026-09-01', 'google_aio', 'found', ['text']),  # retry replaces older result
            (1, '2026-09-01', 'google_aio', 'found', ['marketplace']),
            (0, '2026-09-01', 'chatgpt', 'error', []),
            (0, '2026-09-03', 'google_aio', 'not_found', []),
            (1, '2026-09-03', 'google_aio', 'skipped', []),
            (0, '2026-09-03', 'chatgpt', 'found', ['text']),
        ], 1):
            q = project['queries'][query]
            row = CloudResult(user_id=user_id, device_id='a'*32, local_result_id=i, local_project_id=1,
                project_name=project['name'], brand_name='ESTUS', query_text=q['text'], group_tag=q['group_tag'],
                query_id=q['id'], project_id=project['id'], service=service, scan_date=day, status=status,
                mention_types_json=json.dumps(kinds), sources_json=json.dumps(['https://geosoft-dent.ru/item', 'https://catalog.test/product', 'javascript:alert(1)']),
                evidence_quote='=Untrusted quote', answer_text='Original answer')
            db.add(row); db.flush(); ids.append(row.id)
        db.commit()
    return client, owner, other, devices, project, ids, url


def test_shared_viewer_reads_results_but_cannot_edit_feedback(tmp_path):
    client, owner, viewer, _, project, ids, url = seed_report(tmp_path)
    with make_session_factory(url)[1]() as db:
        viewer_id = db.scalar(select(User.id).where(User.email == 'other@test.example'))
        db.add(ProjectMember(project_id=project['id'], user_id=viewer_id, role='viewer'))
        db.commit()
    base = f"/api/v1/control/projects/{project['id']}/mentions"
    assert client.get(base + '/overview', headers=viewer).status_code == 200
    assert client.get(base + f'/results/{ids[1]}', headers=viewer).status_code == 200
    assert client.put(base + f'/results/{ids[1]}/feedback', headers=viewer,
                      json={'label': 'false_positive'}).status_code == 404


def test_selected_report_dates_share_summary_sources_and_export(tmp_path):
    client, owner, other, _, project, _, url = seed_report(tmp_path)
    base = f"/api/v1/control/projects/{project['id']}/mentions"
    args = {'date_from': '2026-09-01', 'date_to': '2026-09-03',
            'report_dates': '2026-09-03', 'max_dates': 30}
    for suffix in ('', '/overview'):
        response = client.get(base + suffix, headers=owner, params=args)
        assert response.status_code == 200, response.text
        data = response.json()
        assert data['dates'] == ['2026-09-03']
        assert [day['date'] for day in data['timeline']] == ['2026-09-03']
        assert data['summary']['checked'] == 3
        if suffix:
            assert all(item['count'] == 1 for item in data['sources'])
            assert all(row['trend']['current']['checked'] <= 2 for row in data['rows'])
    exported = client.get(base + '/export/mentions', headers=owner, params=args)
    wb = load_workbook(io.BytesIO(exported.content))
    assert all('2026-09-03' in cell.value for cell in list(wb['Упоминаемость'].rows)[0][2:])
    assert client.get(base + '/overview', headers=other, params=args).status_code == 404
    for changes in ({'report_dates': ''}, {'report_dates': '2026-02-30'},
                    {'report_dates': '2026-09-04'}, {'report_dates': ','.join(['2026-09-03'] * 31)},
                    {'max_dates': 31}, {'max_dates': 0}, {'compare': True}, {'date_mode':'monthly'}):
        assert client.get(base + '/overview', headers=owner, params={**args, **changes}).status_code == 422
    limited = {key: value for key, value in args.items() if key != 'report_dates'}
    limited['max_dates'] = 1
    overview = client.get(base + '/overview', headers=owner, params=limited).json()
    matrix = client.get(base, headers=owner, params=limited).json()
    assert overview['dates'] == matrix['dates'] == ['2026-09-03']
    assert overview['summary']['checked'] == matrix['summary']['checked'] == 3
    assert [day['date'] for day in overview['timeline']] == ['2026-09-03']
    assert all(item['count'] == 1 for item in overview['sources'])
    wb = load_workbook(io.BytesIO(client.get(base + '/export/mentions', headers=owner, params=limited).content))
    assert all('2026-09-03' in cell.value for cell in list(wb['Упоминаемость'].rows)[0][2:])
    _, sessions = make_session_factory(url)
    with sessions() as db:
        user_id = db.get(ControlProject, project['id']).user_id
        query = project['queries'][0]
        for day in range(4, 16):
            db.add(CloudResult(user_id=user_id, device_id='a'*32, local_result_id=100+day,
                local_project_id=1, project_name=project['name'], brand_name='ESTUS',
                query_text=query['text'], group_tag=query['group_tag'], query_id=query['id'],
                project_id=project['id'], service='google_aio', scan_date=f'2026-09-{day:02}',
                status='found', mention_types_json='["text"]', sources_json='[]'))
        db.add(CloudResult(user_id=user_id, device_id='a'*32, local_result_id=200,
            local_project_id=1, project_name=project['name'], brand_name='ESTUS',
            query_text='Бренд', group_tag='Legacy', query_id='removed-query',
            project_id=project['id'], service='google_aio', scan_date='2026-09-15',
            status='found', mention_types_json='["text"]', sources_json='[]'))
        db.commit()
    scoped = client.get(base + '/overview', headers=owner, params={
        'date_from':'2026-09-01','date_to':'2026-09-15','services':'chatgpt','max_dates':10}).json()
    assert scoped['dates'] == ['2026-09-01','2026-09-03']
    assert scoped['available_dates'] == ['2026-09-01','2026-09-03']
    assert scoped['summary']['checked'] == 1 and scoped['summary']['issues'] == 1
    monthly = client.get(base + '/overview', headers=owner, params={
        'date_from':'2026-09-01','date_to':'2026-09-15','services':'chatgpt','max_dates':10,'date_mode':'monthly'}).json()
    assert monthly['dates'] == ['2026-09-03']
    assert monthly['summary']['checked'] == 1 and monthly['summary']['issues'] == 0
    empty = client.get(base + '/overview', headers=owner, params={
        'date_from':'2026-09-01','date_to':'2026-09-15','max_dates':10,'date_mode':'monthly','search':'Not checked'})
    assert empty.status_code == 200
    assert empty.json()['dates'] == [] and empty.json()['summary']['checked'] == 0
    legacy = client.get(base + '/overview', headers=owner, params={
        'date_from':'2026-09-01','date_to':'2026-09-15','max_dates':10,'date_mode':'monthly','search':'бренд'}).json()
    assert legacy['available_dates'] == legacy['dates'] == ['2026-09-15']
    assert legacy['summary']['checked'] == 1


def test_calendar_daily_dedup_filters_and_comparison(tmp_path):
    client, owner, other, devices, project, ids, _ = seed_report(tmp_path)
    base = f"/api/v1/control/projects/{project['id']}/mentions"
    args = {'date_from':'2026-09-01','date_to':'2026-09-03'}
    res = client.get(base, headers=owner, params=args)
    assert res.status_code == 200, res.text
    data = res.json()
    assert data['summary'] == {'found':3,'checked':5,'issues':1,'skipped':1,'visibility_pct':60.0}
    assert data['timeline'][1]['visibility_pct'] is None
    assert data['rows'][0]['cells']['2026-09-01']['google_aio']['id'] == ids[1]
    assert data['rows'][2]['cells'] == {}
    assert data['comparison']['delta'] == -66.7
    filtered = client.get(base, headers=owner, params={**args,'include_cards':'false','group':''}).json()
    assert filtered['summary']['found'] == 0 and filtered['summary']['checked'] == 2
    assert filtered['total'] == 1
    assert filtered['rows'][0]['cells']['2026-09-01']['google_aio']['found'] is False
    filtered = client.get(base, headers=owner, params={**args,'service':'chatgpt','search':'Original'}).json()
    assert filtered['summary']['checked'] == 1 and filtered['summary']['issues'] == 1
    assert filtered['services'] == ['chatgpt']
    page = client.get(base, headers=owner, params={**args,'limit':1,'offset':1,'date_limit':1,'date_offset':1}).json()
    assert page['total'] == 3 and len(page['rows']) == 1
    assert page['visible_dates'] == ['2026-09-01']
    assert page['summary'] == data['summary']
    comparison = client.get(base, headers=owner, params={**args,'compare':'true','date_to':'2026-09-02'}).json()
    assert comparison['visible_dates'] == ['2026-09-01','2026-09-02']
    assert comparison['comparison']['delta'] is None
    assert comparison['comparison']['to']['checked'] == 0
    for params in ({'date_from':'2026-09-04','date_to':'2026-09-01'}, {'date_from':'2020-01-01','date_to':'2026-01-01'}, {'service':'invalid'}):
        assert client.get(base, headers=owner, params=params).status_code == 422
    assert client.get(base, headers=other).status_code == 404
    assert client.get(base, headers=devices[0]).status_code == 403
    detail = client.get(f'{base}/results/{ids[1]}', headers=owner).json()
    assert detail['answer_text'] == 'Original answer'
    assert client.get(f'{base}/results/{ids[1]}', headers=other).status_code == 404
    assert client.get(f'{base}/export/mentions', headers=other).status_code == 404


def test_absent_ai_block_counts_as_no_mention_across_reports(tmp_path, monkeypatch):
    client, owner, _, _, project, ids, url = seed_report(tmp_path)
    base = f"/api/v1/control/projects/{project['id']}/mentions"
    args = {'date_from': '2026-09-03', 'date_to': '2026-09-03',
            'service': 'google_aio', 'search': 'Dental'}
    expected = {'found': 0, 'checked': 1, 'issues': 0, 'skipped': 1, 'visibility_pct': 0.0}
    for suffix in ('', '/overview'):
        data = client.get(base + suffix, headers=owner, params=args).json()
        assert data['summary'] == data['timeline'][0]['services']['google_aio'] == expected
        assert data['timeline'][0]['visibility_pct'] == 0
        cell = data['rows'][0]['latest']['google_aio'] if suffix else data['rows'][0]['cells']['2026-09-03']['google_aio']
        assert cell == {'id': ids[5], 'status': 'skipped', 'found': False, **({'date': '2026-09-03'} if suffix else {})}
        if suffix:
            assert data['service_stats']['google_aio']['current'] == expected
            assert data['rows'][0]['trend']['current'] == expected
        for option in ('include_cards', 'include_external'):
            filtered = client.get(base + suffix, headers=owner, params={**args, option: 'false'}).json()
            assert filtered['summary'] == expected
        compared = client.get(base + suffix, headers=owner, params={**args, 'date_from': '2026-09-01', 'compare': 'true'}).json()
        assert compared['comparison']['delta'] == -100
        if suffix:
            assert compared['trend']['delta'] == -100
    monkeypatch.setattr('server.app.utcnow', lambda: datetime(2026, 9, 3, tzinfo=timezone.utc))
    legacy = client.get('/api/v1/reports', headers=owner, params={'project': project['id'], 'days': 3}).json()
    assert legacy['summary']['found'] == 3 and legacy['summary']['checked'] == 5
    assert legacy['summary']['visibility_pct'] == 60
    assert legacy['summary']['by_service']['google_aio'] == {'found': 2, 'checked': 4}
    assert legacy['summary']['by_date']['2026-09-03'] == {'found': 1, 'checked': 3}
    wb = load_workbook(io.BytesIO(client.get(base + '/export/mentions', headers=owner, params=args).content))
    assert wb['Упоминаемость']['C2'].value == 'Нет AI-блока'
    notes = dict(list(wb['О выгрузке'].values)[1:])
    assert 'Отсутствие AI-блока считается проверкой без упоминания' in notes['Упоминаемость']
    # An AI block appearing without the brand must not lower visibility.
    _, sessions = make_session_factory(url)
    with sessions() as db:
        db.get(CloudResult, ids[5]).status = 'not_found'
        db.commit()
    for suffix in ('', '/overview'):
        after = client.get(base + suffix, headers=owner, params=args).json()['summary']
        assert after == {**expected, 'skipped': 0}


@pytest.mark.parametrize('status', ['error', 'captcha', 'auth_required', 'limit_reached'])
def test_technical_issues_are_not_negative_measurements(status):
    from server.reporting import count, empty_stats
    stats = empty_stats()
    count(stats, SimpleNamespace(status=status), True)
    assert stats == {'found': 0, 'checked': 0, 'issues': 1, 'skipped': 0, 'visibility_pct': None}


@pytest.mark.parametrize('source,expected', [
    ('https://geosoft-dent.ru/item', 1), ('https://www.geosoft-dent.ru/item', 1),
    ('https://shop.geosoft-dent.ru/item', 1), ('https://catalog.test/item', 0),
    ('https://geosoft-dent.ru.evil.test/item', 0), ('https://user:pass@geosoft-dent.ru/item', 0),
    ('javascript://geosoft-dent.ru/item', 0), ('', 0),
])
def test_own_source_filter_matches_domains_and_preserves_history(tmp_path, source, expected):
    client, owner, _, _, project, ids, url = seed_report(tmp_path)
    _, sessions = make_session_factory(url)
    with sessions() as db:
        row = db.get(CloudResult, ids[1])
        row.sources_json = json.dumps([source, 'https://catalog.test/other']) if source else '[]'
        db.commit()
    base = f"/api/v1/control/projects/{project['id']}/mentions"
    args = {'date_from': '2026-09-01', 'date_to': '2026-09-01', 'search': 'Original',
            'service': 'google_aio', 'include_external': 'false'}
    for suffix in ('', '/overview'):
        data = client.get(base + suffix, headers=owner, params=args).json()
        assert data['summary']['found'] == expected and data['summary']['checked'] == 1
        assert data['timeline'][0]['found'] == expected
        cell = data['rows'][0]['latest']['google_aio'] if suffix else data['rows'][0]['cells']['2026-09-01']['google_aio']
        assert cell['found'] is bool(expected)
        if suffix:
            assert all(s['own'] for s in data['sources'])
    exported = client.get(base + '/export/mentions', headers=owner, params=args)
    notes = dict(list(load_workbook(io.BytesIO(exported.content))['О выгрузке'].values)[1:])
    assert notes['Учёт внешних источников'] == 'Исключены из упоминаний'
    detail = client.get(f'{base}/results/{ids[1]}', headers=owner).json()
    assert detail['status'] == 'found' and detail['answer_text'] == 'Original answer'
    with sessions() as db:
        p = db.get(ControlProject, project['id'])
        config = json.loads(p.config_json); config['brand_domains'] = []
        p.config_json = json.dumps(config); db.commit()
    assert client.get(base, headers=owner, params=args).json()['summary']['found'] == 0


def test_overview_sources_count_answers_and_keep_intermediate_compare_dates(tmp_path):
    client, owner, other, devices, project, ids, url = seed_report(tmp_path)
    _, sessions = make_session_factory(url)
    with sessions() as db:
        row = db.get(CloudResult, ids[1])
        row.scan_date = '2026-09-02'
        row.sources_json = json.dumps(['https://catalog.test/a', 'https://catalog.test/b',
                                      'https://sub.geosoft-dent.ru/x', 'https://user:pass@bad.test/x',
                                      'javascript:alert(1)'])
        db.commit()
    base = f"/api/v1/control/projects/{project['id']}/mentions/overview"
    args = {'date_from': '2026-09-01', 'date_to': '2026-09-03', 'compare': 'true', 'limit': 1}
    response = client.get(base, headers=owner, params=args)
    assert response.status_code == 200
    data = response.json()
    sources = {s['domain']: s for s in data['sources']}
    assert sources['catalog.test']['count'] == 3
    assert sources['sub.geosoft-dent.ru']['count'] == 1
    assert sources['sub.geosoft-dent.ru']['own'] is True
    assert sources['sub.geosoft-dent.ru']['results'][0]['id'] == ids[1]
    assert not any(s['domain'] == 'bad.test' for s in data['sources'])
    assert len(data['rows']) == 3 and 'measurements' not in data
    assert all(set(q['cells']) <= {'2026-09-01', '2026-09-03'} for q in data['rows'])
    assert response.headers['cache-control'] == 'no-store'
    assert client.get(base, headers=other, params=args).status_code == 404
    assert client.get(base, headers=devices[0], params=args).status_code == 403
    boundary = client.get(base, headers=owner, params={'date_from': '0001-01-01', 'date_to': '0001-01-01'})
    assert boundary.status_code == 200
    assert boundary.json()['trend']['delta'] is None
    export = client.get(base + '/export/sources', headers=owner, params={**args, 'source_kind': 'own', 'source_search': 'sub.'})
    assert export.status_code == 200
    sheet = load_workbook(io.BytesIO(export.content))['Источники']
    assert list(sheet.values)[1:] == [('sub.geosoft-dent.ru', 'Сайт бренда', 1, 1, 'Google AI Overview')]
    notes = dict(list(load_workbook(io.BytesIO(export.content))['О выгрузке'].values)[1:])
    assert notes['Режим'] == 'За период'  # Sources include intermediate dates even from a compared report.
    assert client.get(base + '/export/sources', headers=other, params=args).status_code == 404
    assert client.get(base + '/export/sources', headers=devices[0], params=args).status_code == 403
    assert client.get(base + '/export/sources', headers=owner, params={**args, 'source_kind': 'invalid'}).status_code == 422


def test_idn_and_private_urls_have_consistent_source_exports(tmp_path):
    client, owner, _, _, project, ids, url = seed_report(tmp_path)
    _, sessions = make_session_factory(url)
    with sessions() as db:
        p = db.get(ControlProject, project['id'])
        config = json.loads(p.config_json); config['brand_domains'] = ['пример.рф']
        p.config_json = json.dumps(config)
        db.get(CloudResult, ids[1]).sources_json = json.dumps([
            'https://xn--e1afmkfd.xn--p1ai/page', 'https://user:pass@catalog.test/private',
            'https://external.test/page'])
        db.commit()
    base = f"/api/v1/control/projects/{project['id']}/mentions"
    args = {'date_from': '2026-09-01', 'date_to': '2026-09-01', 'search': 'Original', 'service': 'google_aio'}
    sources = client.get(base + '/overview', headers=owner, params=args).json()['sources']
    assert {s['domain']: s['own'] for s in sources} == {'xn--e1afmkfd.xn--p1ai': True, 'external.test': False}
    exported = client.get(base + '/export/sources', headers=owner, params=args)
    rows = list(load_workbook(io.BytesIO(exported.content))['Внешние источники'].values)[1:]
    assert len(rows) == 1 and rows[0][0] == 'https://external.test/page'


@pytest.mark.parametrize('start,end,compare,baseline_start,baseline_end', [
    ('2026-10-10', '2026-10-12', False, '2026-10-07', '2026-10-09'),
    ('2026-10-10', '2026-10-10', False, '2026-10-09', '2026-10-09'),
    ('2026-01-01', '2026-01-03', False, '2025-12-29', '2025-12-31'),
    ('2024-03-01', '2024-03-01', False, '2024-02-29', '2024-02-29'),
    ('2026-10-10', '2026-10-12', True, '2026-10-10', '2026-10-10'),
])
def test_overview_previous_calendar_periods(tmp_path, start, end, compare, baseline_start, baseline_end):
    client, owner, _, _, project, _, _ = seed_report(tmp_path)
    response = client.get(f"/api/v1/control/projects/{project['id']}/mentions/overview", headers=owner,
                          params={'date_from': start, 'date_to': end, 'compare': str(compare).lower()})
    assert response.status_code == 200
    trend = response.json()['trend']
    assert (trend['baseline_from'], trend['baseline_to']) == (baseline_start, baseline_end)
    assert trend['delta'] is None and trend['relative'] is None and not trend['baseline_complete']


def test_overview_partial_baselines_are_local_to_query_and_service(tmp_path):
    client, owner, _, _, project, _, url = seed_report(tmp_path)
    _, sessions = make_session_factory(url)
    with sessions() as db:
        p = db.get(ControlProject, project['id'])
        query = project['queries'][0]
        for i, (day, service, status) in enumerate([
            ('2026-08-29', 'google_aio', 'not_found'), ('2026-08-30', 'google_aio', 'not_found'),
            ('2026-08-31', 'google_aio', 'not_found'), ('2026-08-29', 'chatgpt', 'found'),
            ('2026-08-30', 'chatgpt', 'found'), ('2026-09-02', 'google_aio', 'found'),
        ], 100):
            db.add(CloudResult(user_id=p.user_id, device_id='a'*32, local_result_id=i, local_project_id=1,
                               project_id=p.id, query_id=query['id'], query_text=query['text'], group_tag='Brand',
                               project_name=p.name, brand_name=p.brand_name, service=service, scan_date=day, status=status))
        db.commit()
    base = f"/api/v1/control/projects/{project['id']}/mentions/overview"
    args = {'date_from':'2026-09-01', 'date_to':'2026-09-03'}
    data = client.get(base, headers=owner, params=args).json()
    assert data['service_stats']['google_aio']['previous']['visibility_pct'] == 0
    assert data['service_stats']['google_aio']['relative'] is None
    assert data['service_stats']['google_aio']['delta'] == 60
    assert data['service_stats']['chatgpt']['delta'] is None
    assert data['rows'][0]['trend']['delta'] == 35
    assert data['rows'][1]['trend']['delta'] is None
    filtered = client.get(base, headers=owner, params={**args,'services':'google_aio','include_cards':'false'}).json()
    assert filtered['trend']['delta'] == 40
    compared = client.get(base, headers=owner, params={**args, 'compare':'true', 'services':'google_aio'}).json()
    assert compared['trend']['current']['visibility_pct'] == 0
    assert compared['trend']['previous']['visibility_pct'] == 100
    assert compared['trend']['delta'] == -100


def test_compact_overview_keeps_all_queries_with_bounded_sql_and_source_references(tmp_path):
    from server.reporting import ReportOptions, overview_report
    client, owner, other, _, project, _, url = seed_report(tmp_path)
    engine, sessions = make_session_factory(url)
    with sessions() as db:
        p = db.get(ControlProject, project['id'])
        queries = json.loads(p.queries_json)
        for i in range(205):
            q = {'id': f'query-{i}', 'text': f'Query {i}', 'group_tag': 'Scale', 'active': True}
            queries.append(q)
            db.add(CloudResult(user_id=p.user_id, device_id='a'*32, local_result_id=100+i, local_project_id=1,
                               project_id=p.id, query_id=q['id'], query_text=q['text'], group_tag='Scale',
                               project_name=p.name, brand_name=p.brand_name, service='google_aio', scan_date='2026-09-02',
                               status='found', answer_text='Private answer never belongs in summary',
                               sources_json=json.dumps([f'https://scale.test/{i}/{j}' for j in range(15)])))
        p.queries_json = json.dumps(queries)
        db.commit()
        statements = []
        def record(*args): statements.append(args[2])
        event.listen(engine, 'before_cursor_execute', record)
        try:
            data = overview_report(db, p, ReportOptions(date_from=date(2026,9,1), date_to=date(2026,9,3), limit=1))
        finally:
            event.remove(engine, 'before_cursor_execute', record)
        # The project reload after commit adds one statement; result columns cause no N+1 reads.
        assert len(statements) <= 9
        assert len(data['rows']) == 208
        assert data['rows'][-1]['cells']['2026-09-02']['google_aio']['found'] is True
        source = next(s for s in data['sources'] if s['domain'] == 'scale.test')
        assert source['count'] == 205 and len(source['results']) == len(source['urls']) == 10
        assert 'Private answer' not in json.dumps(data)
        assert data['rows'][-1]['latest']['google_aio']['date'] == '2026-09-02'
        p.archived = True
        db.commit()
    assert client.get(f"/api/v1/control/projects/{project['id']}/mentions/overview", headers=owner).status_code == 404


def test_overview_history_contains_every_actual_scan_date_and_latest_retry(tmp_path):
    client, owner, other, devices, project, ids, url = seed_report(tmp_path)
    _, sessions = make_session_factory(url)
    with sessions() as db:
        p = db.get(ControlProject, project['id'])
        q = project['queries'][0]
        for i in range(4, 10):
            db.add(CloudResult(user_id=p.user_id, device_id='a'*32, local_result_id=100+i,
                               local_project_id=1, project_id=p.id, query_id=q['id'], query_text=q['text'],
                               project_name=p.name, brand_name=p.brand_name, service='google_aio',
                               scan_date=f'2026-09-{i:02}', status='not_found', answer_text='Private history'))
        db.commit()
    path = f"/api/v1/control/projects/{project['id']}/mentions/overview"
    args = {'date_from':'2026-08-31', 'date_to':'2026-09-10', 'services':'google_aio,chatgpt'}
    data = client.get(path, headers=owner, params=args).json()
    dates = ['2026-09-01', '2026-09-03'] + [f'2026-09-{i:02}' for i in range(4, 10)]
    assert data['dates'] == dates
    assert data['visible_dates'] == dates[-4:]
    assert sorted(data['rows'][0]['cells']) == dates[-4:]
    older = client.get(path, headers=owner, params={**args, 'date_offset':4}).json()
    assert older['visible_dates'] == dates[:4]
    assert sorted(older['rows'][0]['cells']) == dates[:4]
    assert older['summary'] == data['summary']
    assert older['sources'] == data['sources']
    assert older['rows'][0]['trend'] == data['rows'][0]['trend']
    assert older['rows'][0]['cells']['2026-09-01']['google_aio']['id'] == ids[1]
    assert data['rows'][0]['latest']['google_aio']['date'] == dates[-1]
    assert 'Private history' not in json.dumps(data)
    assert client.get(path, headers=other, params=args).status_code == 404
    assert client.get(path, headers=devices[0], params=args).status_code == 403
    compared = client.get(path, headers=owner, params={**args, 'compare':'true'}).json()
    assert compared['dates'] == compared['visible_dates'] == ['2026-08-31', '2026-09-10']
    assert not any(q['cells'] for q in compared['rows'])


def test_exports_use_all_filtered_rows_and_safe_text(tmp_path):
    client, owner, _, _, project, _, _ = seed_report(tmp_path)
    base = f"/api/v1/control/projects/{project['id']}/mentions/export"
    args = {'date_from':'2026-09-01','date_to':'2026-09-03','offset':1,'limit':1,'include_cards':'false'}
    res = client.get(base+'/mentions', headers=owner, params=args)
    assert res.status_code == 200, res.text
    wb = load_workbook(io.BytesIO(res.content))
    ws = wb['Упоминаемость']
    assert ws.max_row == 4 and ws.max_column == 6
    assert ws['A2'].value == '=Original equipment?' and ws['A2'].data_type == 's'
    assert ws['C3'].value == 'Карточки исключены'
    assert ws['C4'].value == 'Не проверялось'
    res = client.get(base+'/sources', headers=owner, params={**args, 'group':'Brand'})
    assert res.status_code == 200, res.text
    wb = load_workbook(io.BytesIO(res.content))
    rows = list(wb['Внешние источники'].values)[1:]
    assert rows and all(row[1] == 'catalog.test' for row in rows)
    assert all(row[5] == 'Brand' for row in rows)
    assert wb['Внешние источники']['G2'].data_type == 's'

def test_multiple_systems_filter_statistics_matrix_and_exports(tmp_path):
    client, owner, _, _, project, _, url = seed_report(tmp_path)
    _, sessions = make_session_factory(url)
    with sessions() as db:
        p = db.get(ControlProject, project['id'])
        config = json.loads(p.config_json)
        config['services'] = ['google_aio', 'chatgpt', 'perplexity', 'alice']
        p.config_json = json.dumps(config)
        for i, service in enumerate(['perplexity', 'alice'], 20):
            q = project['queries'][0]
            db.add(CloudResult(user_id=p.user_id, device_id='a'*32, local_result_id=i, local_project_id=1,
                project_id=p.id, query_id=q['id'], project_name=p.name, brand_name=p.brand_name,
                query_text=q['text'], group_tag='Brand', service=service,
                scan_date='2026-09-02' if service == 'perplexity' else '2026-09-01', status='found',
                sources_json='["https://external.test/page"]'))
        db.commit()
    base = f"/api/v1/control/projects/{project['id']}/mentions"
    args = {'date_from': '2026-09-01', 'date_to': '2026-09-03', 'services': 'google_aio,chatgpt,alice'}
    res = client.get(base, headers=owner, params=args)
    assert res.status_code == 200, res.text
    data = res.json()
    assert data['services'] == ['google_aio', 'chatgpt', 'alice']
    assert data['available_services'] == ['google_aio', 'chatgpt', 'perplexity', 'alice']
    assert data['summary']['checked'] == 6 and data['summary']['found'] == 4
    assert data['summary']['visibility_pct'] == 66.7
    assert data['dates'] == ['2026-09-01', '2026-09-02', '2026-09-03']
    assert data['timeline'][1]['visibility_pct'] is None
    assert all('perplexity' not in day['services'] for day in data['timeline'])
    assert all('perplexity' not in cells for q in data['rows'] for cells in q['cells'].values())
    compared = client.get(base, headers=owner, params={**args, 'compare':'true'}).json()
    assert compared['comparison']['from']['checked'] == 3
    assert compared['comparison']['delta'] == -66.7
    for kind in ('mentions', 'sources'):
        res = client.get(base+'/export/'+kind, headers=owner, params=args)
        assert res.status_code == 200, res.text
        wb = load_workbook(io.BytesIO(res.content))
        assert not any('Perplexity' in str(cell) for ws in wb for row in ws.values for cell in row)
        assert any('Алиса AI' in str(cell) for ws in wb for row in ws.values for cell in row)
        if kind == 'mentions':
            assert any(str(cell).startswith('2026-09-02') for cell in next(wb['Упоминаемость'].values))
    for selected in ('', 'invalid', 'google_aio,unknown', 'google_aio,'):
        assert client.get(base, headers=owner, params={**args,'services':selected}).status_code == 422
    assert client.get(base, headers=owner, params={**args,'service':'chatgpt'}).status_code == 422
    single = client.get(base, headers=owner, params={**args,'services':'alice,alice'}).json()
    assert single['services'] == ['alice'] and single['summary']['checked'] == 1


def test_filtered_service_keeps_empty_selected_date_and_export_label(tmp_path):
    client, owner, _, _, project, _, _ = seed_report(tmp_path)
    base = f"/api/v1/control/projects/{project['id']}/mentions"
    args = {'date_from': '2026-09-02', 'date_to': '2026-09-02', 'services': 'chatgpt'}
    data = client.get(base, headers=owner, params=args).json()
    assert data['dates'] == data['visible_dates'] == ['2026-09-02']
    assert data['services'] == ['chatgpt'] and data['summary']['checked'] == 0
    assert all(row['cells'] == {} for row in data['rows'])
    wb = load_workbook(io.BytesIO(client.get(base + '/export/mentions', headers=owner, params=args).content))
    assert wb['Упоминаемость']['C2'].value == 'Не проверялось'


def test_report_default_period_ends_today(tmp_path, monkeypatch):
    from server import reporting
    monkeypatch.setattr(reporting, 'utcnow', lambda: datetime(2026, 10, 2, 21, 30, tzinfo=timezone.utc))
    client, owner, _, _, project, _, _ = seed_report(tmp_path)
    data = client.get(f"/api/v1/control/projects/{project['id']}/mentions", headers=owner).json()
    assert data['date_to'] == '2026-10-02'
