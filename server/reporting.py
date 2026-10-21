"""Calendar-based mention matrices and exports share the same daily measurements."""
import io
import json
from collections import Counter
from datetime import date, timedelta
from typing import Literal
from urllib.parse import quote, urlsplit

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import load_only

from server.models import CloudResult, ControlProject, utcnow
from server.project_access import project_access

SERVICES = {"google_aio": "Google AI Overview", "chatgpt": "ChatGPT", "perplexity": "Perplexity",
            "alice": "Алиса AI", "yandex_neuro": "Яндекс Поиск с Алисой", "gemini": "Gemini"}
LABELS = {"found": "Упоминание", "not_found": "Нет упоминания", "skipped": "Нет AI-блока",
          "error": "Ошибка", "captcha": "Капча", "auth_required": "Нужен вход", "limit_reached": "Лимит сервиса"}


class ReportOptions(BaseModel):
    date_from: date | None = None
    date_to: date | None = None
    compare: bool = False
    include_cards: bool = True
    include_external: bool = True
    service: str = Field(default="", max_length=40)
    services: str | None = Field(default=None, max_length=200)
    group: str | None = Field(default=None, max_length=120)
    search: str = Field(default="", max_length=2000)
    offset: int = Field(default=0, ge=0)
    limit: int = Field(default=50, ge=1, le=200)
    date_offset: int = Field(default=0, ge=0)
    date_limit: int = Field(default=4, ge=1, le=64)
    max_dates: int | None = Field(default=None, ge=1, le=30)
    report_dates: str | None = Field(default=None, max_length=329)
    date_mode: Literal['period', 'monthly'] = 'period'


def owned_project(db, user_id, project_id):
    return project_access(db, user_id, project_id)[0]


def source_host(source):
    try:
        parsed = urlsplit(source)
        if parsed.scheme in ('http', 'https') and not parsed.username and not parsed.password:
            return (parsed.hostname or '').lower().removeprefix('www.').encode('idna').decode('ascii')
    except (ValueError, UnicodeError):
        pass
    return ''


def brand_hosts(project):
    return {host for domain in json.loads(project.config_json).get('brand_domains', [])
            if (host := source_host(domain if '://' in domain else 'https://' + domain))}


def found(row, cards, own=None, status=None):
    kinds = set(json.loads(row.mention_types_json or "[]"))
    if (status or row.status) != "found" or not (cards or not kinds or bool(kinds - {"marketplace", "card"})):
        return False
    if own is None:
        return True
    return any(host == d or host.endswith('.' + d)
               for source in json.loads(row.sources_json or '[]') if (host := source_host(source)) for d in own)


def empty_stats():
    return {"found": 0, "checked": 0, "issues": 0, "skipped": 0, "visibility_pct": None}


def count(stats, row, cards, own=None, status=None):
    status = status or row.status
    if status in ("found", "not_found", "skipped"):
        stats["checked"] += 1
        stats["found"] += int(found(row, cards, own, status))
        stats["skipped"] += int(status == "skipped")
        stats["visibility_pct"] = round(100 * stats["found"] / stats["checked"], 1)
    else:
        stats["issues"] += 1


def collect_report(db, project, options, *, with_sources=False, compact=False):
    from server.scan_feedback import KEY, effective_status
    corrections = {e["result_id"]: e for e in json.loads(project.config_json).get(KEY, [])}
    own = None if options.include_external else brand_hosts(project)
    scope = (CloudResult.user_id == project.user_id, CloudResult.project_id == project.id)
    end = options.date_to or utcnow().date()
    start = options.date_from or max(date.min, end - timedelta(days=min(29, end.toordinal() - 1)))
    if start > end or (end - start).days > 365:
        raise HTTPException(422, "Выберите период от 1 до 366 дней, начало не позже конца")
    picked = None
    if options.date_mode == 'monthly' and (options.compare or options.report_dates is not None):
        raise HTTPException(422, "Выберите один режим дат отчёта")
    if options.report_dates is not None:
        if options.compare:
            raise HTTPException(422, "Выберите сравнение двух дат или произвольные даты")
        try:
            values = options.report_dates.split(',')
            picked = sorted({date.fromisoformat(value).isoformat() for value in values})
        except ValueError as exc:
            raise HTTPException(422, "Выберите корректные даты отчёта") from exc
        if not 1 <= len(values) <= (options.max_dates or 30) or any(day < start.isoformat() or day > end.isoformat() for day in picked):
            raise HTTPException(422, "Выберите до 30 дат внутри периода отчёта")
    if options.service and options.service not in SERVICES:
        raise HTTPException(422, "Неизвестная ИИ-система")
    selected = None
    if options.services is not None:
        selected = list(dict.fromkeys(s.strip() for s in options.services.split(",")))
        if not selected or any(s not in SERVICES for s in selected):
            raise HTTPException(422, "Выберите хотя бы одну известную ИИ-систему")
        if options.service:
            raise HTTPException(422, "Используйте либо service, либо services")
    elif options.service:
        selected = [options.service]
    configured = json.loads(project.config_json).get("services", [])
    observed = set(db.scalars(select(CloudResult.service).where(*scope).distinct()))
    available_services = [s for s in SERVICES if s in set(configured) | observed]
    queries = {q["id"]: {"id": q["id"], "text": q["text"], "group_tag": q.get("group_tag", ""), "cells": {}}
               for q in json.loads(project.queries_json)}
    by_text = {q["text"]: q["id"] for q in queries.values()}
    date_scope = [*scope, CloudResult.service.in_(selected if selected is not None else SERVICES)]
    if options.group is not None or options.search:
        matching_dates = set()
        metadata = db.execute(select(CloudResult.scan_date, CloudResult.query_id,
            CloudResult.query_text, CloudResult.group_tag).where(*date_scope).distinct())
        for day, query_id, text, group_tag in metadata:
            query = queries.get(query_id) or queries.get(by_text.get(text)) or {'text': text, 'group_tag': group_tag or ''}
            if (options.group is None or query['group_tag'] == options.group) and options.search.casefold() in query['text'].casefold():
                matching_dates.add(day)
        available_dates = sorted(matching_dates)
    else:
        available_dates = list(db.scalars(select(CloudResult.scan_date).where(*date_scope).distinct()
                                          .order_by(CloudResult.scan_date)))
    columns = [
        CloudResult.id, CloudResult.query_id, CloudResult.query_text, CloudResult.group_tag,
        CloudResult.service, CloudResult.scan_date, CloudResult.status, CloudResult.mention_types_json,
    ]
    if with_sources:
        columns += [CloudResult.sources_json, CloudResult.evidence_quote]
    elif not options.include_external:
        columns += [CloudResult.sources_json]
    measurements = db.scalars(select(CloudResult).options(load_only(*columns)).where(
        *scope, CloudResult.scan_date >= start.isoformat(), CloudResult.scan_date <= end.isoformat(),
        CloudResult.service.in_(selected if selected is not None else SERVICES))
      .order_by(CloudResult.id.desc())).all()
    measurement_dates = set(db.scalars(select(CloudResult.scan_date).where(
        *scope, CloudResult.scan_date >= start.isoformat(), CloudResult.scan_date <= end.isoformat()).distinct()))
    daily = {}
    for row in measurements:
        key = row.query_id if row.query_id in queries else by_text.get(row.query_text, row.query_id or "text:" + row.query_text)
        queries.setdefault(key, {"id": key, "text": row.query_text, "group_tag": row.group_tag or "", "cells": {}})
        daily.setdefault((key, row.service, row.scan_date), row)
    group_names = sorted({q["group_tag"] for q in queries.values() if q["group_tag"]}, key=str.casefold)
    queries = {key: q for key, q in queries.items()
               if (options.group is None or q["group_tag"] == options.group)
               and options.search.casefold() in q["text"].casefold()}
    daily = {key: row for key, row in daily.items() if key[0] in queries}
    service_ids = [s for s in SERVICES if s in selected] if selected is not None else available_services
    table_dates = sorted({start.isoformat(), end.isoformat()}) if options.compare else sorted(
        {key[2] for key in daily} if options.max_dates is not None else
        measurement_dates if compact else {start.isoformat(), end.isoformat()} | measurement_dates)
    if picked is not None:
        table_dates = picked
        daily = {key: row for key, row in daily.items() if key[2] in picked}
    if options.date_mode == 'monthly':
        last_by_month = {}
        for day in sorted({key[2] for key in daily}):
            last_by_month[day[:7]] = day
        table_dates = list(last_by_month.values())[-(options.max_dates or 30):]
    if options.max_dates is not None and not options.compare:
        table_dates = table_dates[-options.max_dates:]
    if picked is not None or options.max_dates is not None or options.date_mode == 'monthly':
        daily = {key: row for key, row in daily.items() if key[2] in table_dates}
    # Page dates from the most recent measurements, then present each window chronologically.
    window = table_dates if options.compare else sorted(list(reversed(table_dates))[options.date_offset:options.date_offset + options.date_limit])
    summary = empty_stats()
    by_date = {}
    for i in range((end - start).days + 1):
        day = (start + timedelta(days=i)).isoformat()
        by_date[day] = {**empty_stats(), "services": {s: empty_stats() for s in service_ids}}
    selected_measurements = []
    for (key, service_id, day), row in daily.items():
        status = effective_status(row, corrections.get(row.id))
        count(by_date[day], row, options.include_cards, own, status)
        count(by_date[day]["services"][service_id], row, options.include_cards, own, status)
        if compact or not options.compare or day in table_dates:
            count(summary, row, options.include_cards, own, status)
            selected_measurements.append((queries[key], row))
        cell = {"id": row.id, "status": status, "found": bool(found(row, options.include_cards, own, status))}
        if compact:
            latest = queries[key].setdefault("latest", {})
            if service_id not in latest or day > latest[service_id]["date"]:
                latest[service_id] = {**cell, "date": day}
        if not compact or with_sources and day in window:
            queries[key]["cells"].setdefault(day, {})[service_id] = cell
    rows = list(queries.values())
    comparison = {"from": by_date[start.isoformat()], "to": by_date[end.isoformat()]}
    left, right = comparison["from"]["visibility_pct"], comparison["to"]["visibility_pct"]
    comparison["delta"] = round(right - left, 1) if left is not None and right is not None else None
    return {"date_from": start.isoformat(), "date_to": end.isoformat(), "available_dates": available_dates,
            "dates": table_dates, "visible_dates": window, "services": service_ids,
            "available_services": available_services, "groups": group_names,
            "summary": summary, "comparison": comparison,
            "timeline": [{"date": d, **stats} for d, stats in by_date.items()
                         if picked is None and options.max_dates is None and options.date_mode == 'period' or d in table_dates],
            "total": len(rows), "rows": rows, "measurements": selected_measurements,
            "filters": {"include_cards": options.include_cards, "include_external": options.include_external, "group": options.group,
                        "service": options.service, "services": service_ids,
                        "search": options.search, "compare": options.compare,
                        "report_dates": picked, "max_dates": options.max_dates, "date_mode": options.date_mode}}


def overview_report(db, project, options):
    """Compact, unpaginated period totals; retain the existing matrix/export contract."""
    from server.scan_feedback import KEY, effective_status
    corrections = {e["result_id"]: e for e in json.loads(project.config_json).get(KEY, [])}
    report = collect_report(db, project, options, with_sources=True, compact=True)
    start, end = date.fromisoformat(report['date_from']), date.fromisoformat(report['date_to'])
    length = (end - start).days + 1
    selected_days = report['dates'] if options.report_dates is not None or options.max_dates is not None or options.date_mode == 'monthly' else None
    has_baseline = (options.compare or start.toordinal() > length) and (selected_days is None or bool(selected_days))
    baseline_end = start if options.compare else start - timedelta(days=1) if has_baseline else None
    baseline_start = start if options.compare else start - timedelta(days=length) if has_baseline else None
    baseline_dates = None
    if selected_days is not None and not options.compare and has_baseline:
        baseline_dates = ','.join((date.fromisoformat(day) - timedelta(days=length)).isoformat()
                                  for day in selected_days)
    baseline = collect_report(db, project, options.model_copy(update={
        'date_from': baseline_start, 'date_to': baseline_end, 'compare': False,
        'report_dates': baseline_dates, 'date_mode': 'period'}), compact=True) if has_baseline else {'measurements': []}
    own = brand_hosts(project)
    allowed = None if options.include_external else own

    def totals(data, only_date=None):
        summary, services, queries, days = empty_stats(), {}, {}, {}
        for query, row in data['measurements']:
            if only_date and row.scan_date != only_date:
                continue
            status = effective_status(row, corrections.get(row.id))
            count(summary, row, options.include_cards, allowed, status)
            count(services.setdefault(row.service, empty_stats()), row, options.include_cards, allowed, status)
            count(queries.setdefault(query['id'], empty_stats()), row, options.include_cards, allowed, status)
            days.setdefault(query['id'], set()).add(row.scan_date)
        return summary, services, queries, days

    current, services, queries, _ = totals(report, end.isoformat() if options.compare else None)
    previous, previous_services, previous_queries, previous_days = totals(baseline)
    expected_days = 1 if options.compare else len(selected_days) if selected_days is not None else length

    def trend(now, before, complete):
        left, right = before['visibility_pct'], now['visibility_pct']
        delta = round(right - left, 1) if complete and left is not None and right is not None else None
        return {'current': now, 'previous': before, 'delta': delta,
                'relative': round(100 * (right - left) / left, 1) if delta is not None and left else None,
                'baseline_complete': complete, 'baseline_from': baseline_start.isoformat() if baseline_start else None,
                'baseline_to': baseline_end.isoformat() if baseline_end else None}

    scan_dates = sorted({row.scan_date for _, row in report['measurements']})
    first_date, last_date = (scan_dates[0], scan_dates[-1]) if scan_dates else (None, None)
    first, _, first_queries, _ = totals(report, first_date) if first_date else (empty_stats(), {}, {}, {})
    last, _, last_queries, _ = totals(report, last_date) if last_date else (empty_stats(), {}, {}, {})

    def period_trend(now, before):
        return {**trend(now, before, len(scan_dates) > 1),
                'baseline_from': first_date, 'baseline_to': first_date}

    report['period_from'], report['period_to'] = first_date, last_date
    report['period_trend'] = period_trend(last, first)

    complete = len({r.scan_date for _, r in baseline['measurements']}) == expected_days
    report['trend'] = trend(current, previous, complete)
    report['summary'] = current
    report['service_stats'] = {s: trend(services.get(s, empty_stats()), previous_services.get(s, empty_stats()),
        len({r.scan_date for _, r in baseline['measurements'] if r.service == s}) == expected_days)
        for s in report['services']}
    for query in report['rows']:
        query['period_trend'] = period_trend(last_queries.get(query['id'], empty_stats()),
                                            first_queries.get(query['id'], empty_stats()))
        query['trend'] = trend(queries.get(query['id'], empty_stats()), previous_queries.get(query['id'], empty_stats()),
                              len(previous_days.get(query['id'], set())) == expected_days)
        query.setdefault('latest', {})
    sources = {}
    for query, row in report.pop('measurements'):
        if not found(row, options.include_cards, allowed, effective_status(row, corrections.get(row.id))):
            continue
        for source in sorted(set(json.loads(row.sources_json or '[]'))):
            host = source_host(source)
            is_own = any(host == d or host.endswith('.' + d) for d in own)
            if not host or not options.include_external and not is_own:
                continue
            item = sources.setdefault(host, {'domain': host, 'results': {}, 'urls': set(), 'queries': set(),
                'services': set(), 'own': is_own})
            item['results'][row.id] = {'id': row.id, 'query_text': query['text'], 'date': row.scan_date, 'service': row.service}
            if len(item['urls']) < 10:
                item['urls'].add(source)
            item['queries'].add(query['id']); item['services'].add(row.service)
    report['sources'] = [{**item, 'count': len(item['results']), 'urls': sorted(item['urls']),
                          'results': sorted(item['results'].values(), key=lambda r: (r['date'], r['id']), reverse=True)[:10],
                          'queries': len(item['queries']), 'services': sorted(item['services'])}
                         for item in sorted(sources.values(), key=lambda item: (-len(item['results']), item['domain']))]
    return report


def workbook(report, project, kind):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    def sheet(title, headers, widths):
        ws = wb.create_sheet(title)
        ws.freeze_panes = "C2" if kind == "mentions" else "A2"
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        header = []
        for value in headers:
            cell = WriteOnlyCell(ws, value=ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="493266")
            header.append(cell)
        ws.append(header)
        return ws
    def append(ws, values):
        # Prompts and external pages are untrusted: export all strings as text, never formulas.
        cells = []
        for value in values:
            cell = WriteOnlyCell(ws, value=ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value)
            if isinstance(value, str):
                cell.data_type = "s"
            cells.append(cell)
        ws.append(cells)
    if kind == "mentions":
        columns = [(d, s) for d in report["dates"] for s in report["services"]]
        if len(columns) * len(report["rows"]) > 2_000_000:
            raise HTTPException(422, "Выгрузка слишком большая. Сократите период или выберите группу.")
        ws = sheet("Упоминаемость", ["Запрос", "Группа"] + [f"{d} · {SERVICES[s]}" for d, s in columns],
                   [65, 25] + [28] * len(columns))
        for q in report["rows"]:
            values = [q["text"], q["group_tag"]]
            for day, service in columns:
                cell = q["cells"].get(day, {}).get(service)
                excluded = "Карточки исключены" if report['filters']['include_external'] else "Не учитывается по фильтрам"
                values.append("Не проверялось" if cell is None else excluded if cell["status"] == "found" and not cell["found"]
                              else LABELS.get(cell["status"], cell["status"]))
            append(ws, values)
    elif kind == "source_summary":
        ws = sheet("Источники", ["Домен", "Тип", "Ответов со ссылкой", "Запросы", "ИИ-системы"], [45, 24, 24, 15, 55])
        for item in report['sources']:
            append(ws, [item['domain'], "Сайт бренда" if item['own'] else "Внешний источник", item['count'],
                        item['queries'], ', '.join(SERVICES.get(s, s) for s in item['services'])])
    else:
        own = brand_hosts(project)
        ws = sheet("Внешние источники", ["Источник", "Домен", "Дата", "ИИ-система", "Запрос", "Группа", "Контекст в ответе ИИ"],
                   [65, 28, 14, 24, 65, 25, 90])
        counts = Counter()
        for query, row in report["measurements"]:
            if not report['filters']['include_external']:
                continue
            for source in dict.fromkeys(json.loads(row.sources_json or "[]")):
                host = source_host(source)
                if not host or any(
                    host == d or host.endswith("." + d) for d in own if d):
                    continue
                counts[host] += 1
                append(ws, [source, host, row.scan_date, SERVICES.get(row.service, row.service),
                            query["text"], query["group_tag"], row.evidence_quote or ""])
        stats = sheet("Домены", ["Домен", "Ссылок в ответах"], [45, 24])
        for host, number in counts.most_common():
            append(stats, [host, number])
    notes = sheet("О выгрузке", ["Параметр", "Значение"], [30, 100])
    append(notes, ["Проект", project.name])
    append(notes, ["Период", report["date_from"] + " — " + report["date_to"]])
    filters = report["filters"]
    append(notes, ["Режим", "Сравнение двух дат" if filters["compare"] and kind != 'source_summary' else "За период"])
    if filters.get('report_dates') is not None:
        append(notes, ['Выбранные даты', ', '.join(filters['report_dates'])])
    if filters.get('max_dates') is not None:
        append(notes, ['Максимально дат', filters['max_dates']])
    if filters.get('date_mode') == 'monthly':
        append(notes, ['Выбор дат', 'Последняя проверка месяца'])
    append(notes, ["Карточки товаров", "Учитываются" if filters["include_cards"] else "Исключены из упоминаний"])
    append(notes, ["Учёт внешних источников", "Учитываются" if filters["include_external"] else "Исключены из упоминаний"])
    append(notes, ["Группа", "Все группы" if filters["group"] is None else filters["group"] or "Без группы"])
    append(notes, ["ИИ-системы", ", ".join(SERVICES[s] for s in filters["services"])])
    append(notes, ["Поиск", filters["search"]])
    append(notes, ["Упоминаемость", "Доля проверок с упоминанием бренда. Отсутствие AI-блока считается проверкой без упоминания. Технические ошибки исключены из расчёта."])
    append(notes, ["Внешние источники", "Сохранённые ссылки из ответов ИИ, без доменов бренда. Контекст — фрагмент ответа ИИ, не цитата с внешнего сайта."])
    if kind == 'source_summary':
        append(notes, ['Источники', 'Только ответы с упоминанием бренда. Несколько ссылок на домен из одного ответа считаются один раз.'])
        append(notes, ['Тип источника', report['source_kind']])
        append(notes, ['Поиск домена', report['source_search']])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def register_reports(app, db_session, current_user):
    router = APIRouter(prefix="/api/v1/control/projects/{project_id}/mentions")

    @router.get("")
    def report(project_id: str, options: ReportOptions = Depends(), user=Depends(current_user), db=Depends(db_session)):
        data = collect_report(db, owned_project(db, user.id, project_id), options)
        data.pop("measurements")
        data["rows"] = [{**q, "cells": {d: cells for d, cells in q["cells"].items() if d in data["visible_dates"]}}
                        for q in data["rows"][options.offset:options.offset + options.limit]]
        return data

    @router.get("/overview")
    def overview(project_id: str, response: Response, options: ReportOptions = Depends(),
                 user=Depends(current_user), db=Depends(db_session)):
        response.headers['Cache-Control'] = 'no-store'
        return overview_report(db, owned_project(db, user.id, project_id), options)

    @router.get("/overview/export/sources")
    def source_summary_export(project_id: str, options: ReportOptions = Depends(),
                              source_kind: Literal['all', 'own', 'external'] = 'all',
                              source_search: str = Query(default='', max_length=254),
                              user=Depends(current_user), db=Depends(db_session)):
        project = owned_project(db, user.id, project_id)
        data = overview_report(db, project, options)
        data['sources'] = [s for s in data['sources'] if
                           (source_kind == 'all' or s['own'] == (source_kind == 'own'))
                           and source_search.casefold() in s['domain'].casefold()]
        data.update(source_kind=source_kind, source_search=source_search)
        filename = f"sources-{project.id[:8]}-{data['date_from']}-{data['date_to']}.xlsx"
        return Response(workbook(data, project, 'source_summary'),
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={'Content-Disposition': f"attachment; filename*=UTF-8''{quote(filename)}", 'Cache-Control': 'no-store'})

    @router.get("/results/{result_id}")
    def detail(project_id: str, result_id: int, response: Response,
               user=Depends(current_user), db=Depends(db_session)):
        project = owned_project(db, user.id, project_id)
        row = db.scalar(select(CloudResult).where(CloudResult.id == result_id,
            CloudResult.user_id == project.user_id, CloudResult.project_id == project_id))
        if not row:
            raise HTTPException(404, "Результат не найден")
        from server.scan_feedback import result_detail
        response.headers["Cache-Control"] = "no-store"
        return result_detail(db, row, shared=user.id != project.user_id)

    @router.get("/export/{kind}")
    def export(project_id: str, kind: Literal["mentions", "sources"], options: ReportOptions = Depends(),
               user=Depends(current_user), db=Depends(db_session)):
        project = owned_project(db, user.id, project_id)
        data = collect_report(db, project, options, with_sources=kind == "sources")
        filename = f"{kind}-{project.id[:8]}-{data['date_from']}-{data['date_to']}.xlsx"
        return Response(workbook(data, project, kind), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}", "Cache-Control": "no-store"})

    app.include_router(router)
